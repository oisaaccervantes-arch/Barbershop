from datetime import date, datetime, timedelta
from io import BytesIO

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from sqlalchemy import select
from sqlalchemy.orm import Session, joinedload

from app.models.attendance import AttendanceRecord, WorkSchedule
from app.models.barber import Barber
from app.models.cash_shift import CashShift
from app.models.receptionist import Receptionist


DAYS = ["Sábado", "Domingo", "Lunes", "Martes", "Miércoles", "Jueves", "Viernes"]
SHIFT_LABELS = {"MORNING": "Matutino", "EVENING": "Vespertino"}
STATUS_LABELS = {"PENDING": "Pendiente", "PRESENT": "Asistió", "ABSENT": "Falta", "REST": "Descanso", "PERMISSION": "Permiso"}
DARK = "2B2118"
GOLD = "E5B94F"
GOLD_SOFT = "F8E4A4"
CREAM = "FCF8F0"
LINE = "D9CEBE"


def employee_name(db: Session, person_type: str, person_id: int) -> str:
    model = Barber if person_type == "BARBER" else Receptionist
    employee = db.get(model, person_id)
    return employee.name if employee else f"Personal #{person_id}"


def local_naive(value: datetime | None) -> datetime | None:
    if value is None:
        return None
    return value.astimezone().replace(tzinfo=None) if value.tzinfo else value


def base_workbook(title: str, subtitle: str):
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Resumen semanal"
    sheet.sheet_view.showGridLines = False
    sheet.merge_cells("A1:H1")
    sheet["A1"] = title
    sheet["A1"].font = Font(size=18, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill("solid", fgColor=DARK)
    sheet["A1"].alignment = Alignment(vertical="center")
    sheet.row_dimensions[1].height = 32
    sheet.merge_cells("A2:H2")
    sheet["A2"] = subtitle
    sheet["A2"].font = Font(size=11, color="66584A")
    sheet["A2"].fill = PatternFill("solid", fgColor=CREAM)
    return workbook, sheet


def style_table(sheet, max_row: int, max_col: int, header_row: int = 4):
    thin = Side(style="thin", color=LINE)
    for cell in sheet[header_row]:
        if cell.column <= max_col:
            cell.fill = PatternFill("solid", fgColor=GOLD)
            cell.font = Font(bold=True, color=DARK)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in sheet.iter_rows(min_row=header_row + 1, max_row=max_row, min_col=1, max_col=max_col):
        for cell in row:
            cell.border = Border(bottom=thin)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    sheet.freeze_panes = f"B{header_row + 1}"
    sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(max_col)}{max_row}"


def workbook_bytes(workbook: Workbook) -> bytes:
    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def export_schedule(db: Session, week_start: date, shift_type: str) -> bytes:
    rows = db.scalars(
        select(WorkSchedule).where(
            WorkSchedule.week_start == week_start,
            WorkSchedule.shift_type == shift_type,
        ).order_by(WorkSchedule.person_type.desc(), WorkSchedule.person_id, WorkSchedule.day_of_week)
    ).all()
    end = week_start + timedelta(days=6)
    workbook, sheet = base_workbook(
        "Bizantino Barbería · Horario semanal",
        f"{week_start:%d/%m/%Y} al {end:%d/%m/%Y} · {SHIFT_LABELS.get(shift_type, shift_type)}",
    )
    sheet.append([])
    sheet.append(["Personal", *[f"{day}\n{(week_start + timedelta(days=index)):%d/%m}" for index, day in enumerate(DAYS)]])
    people: dict[tuple[str, int], dict[int, WorkSchedule]] = {}
    for row in rows:
        people.setdefault((row.person_type, row.person_id), {})[row.day_of_week] = row
    for (person_type, person_id), days in people.items():
        values = [f"{employee_name(db, person_type, person_id)}\n{'Barbero' if person_type == 'BARBER' else 'Recepción'}"]
        for day in range(7):
            schedule = days.get(day)
            if schedule is None:
                values.append("—")
            elif schedule.status == "REST":
                values.append("DESCANSO")
            else:
                meal = f"\nComida: {schedule.meal_start:%H:%M}–{schedule.meal_end:%H:%M}" if schedule.meal_start and schedule.meal_end else ""
                values.append(f"Entrada: {schedule.start_time:%H:%M}\nSalida: {schedule.end_time:%H:%M}{meal}")
        sheet.append(values)
    style_table(sheet, max(sheet.max_row, 4), 8)
    sheet.column_dimensions["A"].width = 24
    for column in range(2, 9):
        sheet.column_dimensions[get_column_letter(column)].width = 22
    for row in range(5, sheet.max_row + 1):
        sheet.row_dimensions[row].height = 58
    return workbook_bytes(workbook)


def export_attendance(db: Session, week_start: date) -> bytes:
    week_end = week_start + timedelta(days=6)
    records = db.scalars(
        select(AttendanceRecord)
        .join(AttendanceRecord.shift)
        .options(joinedload(AttendanceRecord.shift))
        .where(
            CashShift.business_date.between(week_start, week_end),
            CashShift.status == "CLOSED",
            (AttendanceRecord.notes.is_(None))
            | (AttendanceRecord.notes != "Turno anterior a la implementación del checador"),
        )
        .order_by(AttendanceRecord.person_name, CashShift.business_date, CashShift.shift_type)
    ).all()
    workbook, summary = base_workbook(
        "Bizantino Barbería · Asistencias",
        f"{week_start:%d/%m/%Y} al {week_end:%d/%m/%Y}",
    )
    summary.append([])
    summary.append(["Personal", *[f"{day}\n{(week_start + timedelta(days=index)):%d/%m}" for index, day in enumerate(DAYS)]])
    people: dict[tuple[str, int, str], dict[int, AttendanceRecord]] = {}
    for record in records:
        key = (record.person_type, record.person_id, record.shift.shift_type)
        day_index = (record.shift.business_date.weekday() - 5) % 7
        people.setdefault(key, {})[day_index] = record
    for (person_type, _person_id, shift_type), days in people.items():
        first = next(iter(days.values()))
        values = [f"{first.person_name}\n{'Barbero' if person_type == 'BARBER' else 'Recepción'} · {SHIFT_LABELS.get(shift_type, shift_type)}"]
        for day in range(7):
            record = days.get(day)
            if record is None:
                values.append("—")
                continue
            programmed = record.scheduled_start.strftime("%H:%M") if record.scheduled_start else "Sin horario"
            arrived = local_naive(record.clock_in).strftime("%H:%M") if record.clock_in else "—"
            if record.status == "PRESENT" and record.scheduled_start and record.clock_in:
                expected = datetime.combine(record.shift.business_date, record.scheduled_start)
                difference = round((local_naive(record.clock_in) - expected).total_seconds() / 60)
                comparison = "A la hora" if difference == 0 else (f"{abs(difference)} min antes" if difference < 0 else f"+{difference} min")
            else:
                comparison = STATUS_LABELS.get(record.status, record.status)
            continuation = "\nContinúa en vespertino" if record.continues_next_shift else ""
            values.append(f"Programado: {programmed}\nLlegó: {arrived}\n{comparison}{continuation}")
        summary.append(values)
    style_table(summary, max(summary.max_row, 4), 8)
    summary.column_dimensions["A"].width = 27
    for column in range(2, 9):
        summary.column_dimensions[get_column_letter(column)].width = 22
    for row in range(5, summary.max_row + 1):
        summary.row_dimensions[row].height = 58

    detail = workbook.create_sheet("Detalle")
    detail.sheet_view.showGridLines = False
    detail.append(["Fecha", "Empleado", "Puesto", "Turno", "Entrada programada", "Entrada real", "Diferencia (min)", "Salida a comida", "Regreso de comida", "Salida programada", "Salida real", "Estado", "Notas"])
    for record in records:
        expected_in = datetime.combine(record.shift.business_date, record.scheduled_start) if record.scheduled_start else None
        expected_out = datetime.combine(record.shift.business_date, record.scheduled_end) if record.scheduled_end else None
        row = detail.max_row + 1
        detail.append([
            record.shift.business_date, record.person_name, "Barbero" if record.person_type == "BARBER" else "Recepción",
            SHIFT_LABELS.get(record.shift.shift_type, record.shift.shift_type), expected_in, local_naive(record.clock_in),
            None, local_naive(record.meal_out), local_naive(record.meal_in), expected_out, local_naive(record.clock_out),
            STATUS_LABELS.get(record.status, record.status), record.notes,
        ])
        detail.cell(row, 7, f'=IF(OR(E{row}="",F{row}=""),"",ROUND((F{row}-E{row})*1440,0))')
    style_table(detail, max(detail.max_row, 1), 13, header_row=1)
    detail.column_dimensions["A"].width = 13
    for column in ["B", "C", "D", "L"]:
        detail.column_dimensions[column].width = 18
    for column in ["E", "F", "H", "I", "J", "K"]:
        detail.column_dimensions[column].width = 20
    detail.column_dimensions["G"].width = 17
    detail.column_dimensions["M"].width = 32
    for cell in detail["A"][1:]:
        cell.number_format = "dd/mm/yyyy"
    for column in ["E", "F", "H", "I", "J", "K"]:
        for cell in detail[column][1:]:
            cell.number_format = "dd/mm/yyyy hh:mm"
    return workbook_bytes(workbook)
